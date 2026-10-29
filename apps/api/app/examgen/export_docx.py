"""EB-404 — Xuất đề ra Word (.docx): bản học sinh / bản giáo viên / ma trận đề.

Dữ liệu lấy từ snapshot của từng mã đề (exam_items.snapshot_json) nên file xuất ra
khớp tuyệt đối với bản xem trước trên web (EB-405). Định dạng <b>/<u>/<i> trong
thân câu và phương án được giữ khi chuyển sang Word.
"""
import io

from bs4 import BeautifulSoup
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from sqlalchemy.orm import Session

from app.models import Exam, ExamItem, ExamVersion, TemplateSection


def _add_rich(paragraph, html_fragment: str, *, base_bold=False, base_italic=False):
    """Chuyển HTML đã sanitize (chỉ còn b/i/u/br) thành các run Word giữ định dạng."""
    soup = BeautifulSoup(html_fragment or "", "html.parser")
    body = soup.body or soup

    def walk(node, bold, italic, underline):
        if isinstance(node, str):
            if node:
                run = paragraph.add_run(node)
                run.bold = bold or base_bold
                run.italic = italic or base_italic
                run.underline = underline
            return
        name = getattr(node, "name", None)
        if name == "br":
            paragraph.add_run().add_break()
            return
        if name in ("b", "strong"):
            bold = True
        elif name in ("i", "em"):
            italic = True
        elif name == "u":
            underline = True
        for child in getattr(node, "children", []):
            walk(child, bold, italic, underline)

    for child in body.children:
        walk(child, False, False, False)


def _version_items(db: Session, version: ExamVersion) -> list[dict]:
    rows = (db.query(ExamItem).filter_by(exam_version_id=version.id)
            .order_by(ExamItem.position).all())
    return [{**r.snapshot_json, "position": r.position,
             "correct_label": r.shuffled_correct_label} for r in rows]


def _header(doc: Document, exam: Exam, version: ExamVersion, teacher: bool):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ĐỀ KIỂM TRA — MÔN: TIẾNG ANH")
    r.bold = True
    r.font.size = Pt(14)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{exam.title}   ·   Mã đề: {version.code}"
                  f"   ·   Thời gian: {exam.duration_minutes or ''} phút")
    r.bold = True
    if teacher:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("(BẢN GIÁO VIÊN — đáp án đúng được in đậm, gạch chân)")
        r.italic = True
    p = doc.add_paragraph("Họ và tên học sinh: ............................................................   Lớp: .................")
    doc.add_paragraph("")


def _render_version(doc: Document, exam: Exam, version: ExamVersion,
                    items: list[dict], teacher: bool):
    _header(doc, exam, version, teacher)
    printed_passages: set[str] = set()
    for it in items:
        if it.get("passage_id") and it["passage_id"] not in printed_passages:
            printed_passages.add(it["passage_id"])
            passage = it.get("passage") or {}
            if passage.get("title"):
                pi = doc.add_paragraph()
                _add_rich(pi, passage["title"], base_italic=True)
            for block in (passage.get("content_html") or "").split("</p>"):
                block = block.strip()
                if block:
                    pp = doc.add_paragraph()
                    pp.paragraph_format.left_indent = Pt(18)
                    _add_rich(pp, block + "</p>" if not block.endswith("</p>") else block,
                              base_italic=True)
        p = doc.add_paragraph()
        r = p.add_run(f"Câu {it['position']}: ")
        r.bold = True
        _add_rich(p, it["stem_html"])
        for o in it["options"]:
            po = doc.add_paragraph()
            po.paragraph_format.left_indent = Pt(24)
            mark = teacher and o["label"] == it["correct_label"]
            ro = po.add_run(f"{o['label']}. ")
            ro.bold = mark
            ro.underline = mark
            _add_rich(po, o["html"], base_bold=mark)
        if teacher and it.get("explanation_html"):
            pe = doc.add_paragraph()
            pe.paragraph_format.left_indent = Pt(24)
            r = pe.add_run("Giải thích: ")
            r.italic = True
            _add_rich(pe, it["explanation_html"], base_italic=True)
    # Bảng đáp án cho giáo viên
    if teacher:
        doc.add_paragraph("")
        p = doc.add_paragraph()
        r = p.add_run(f"BẢNG ĐÁP ÁN — Mã đề {version.code}")
        r.bold = True
        table = doc.add_table(rows=0, cols=10)
        table.style = "Table Grid"
        for start in range(0, len(items), 10):
            chunk = items[start:start + 10]
            row = table.add_row().cells
            for j, it in enumerate(chunk):
                row[j].text = str(it["position"])
            row = table.add_row().cells
            for j, it in enumerate(chunk):
                row[j].text = it["correct_label"] or ""


def build_docx(db: Session, exam: Exam, *, kind: str,
               version_code: str | None = None) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)
    versions = (db.query(ExamVersion).filter_by(exam_id=exam.id)
                .order_by(ExamVersion.code).all())
    if version_code:
        versions = [v for v in versions if v.code == version_code]
    if kind == "matrix":
        _render_matrix(doc, db, exam, versions)
    else:
        teacher = kind == "teacher"
        for idx, v in enumerate(versions):
            if idx > 0:
                doc.add_page_break()
            _render_version(doc, exam, v, _version_items(db, v), teacher)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _render_matrix(doc: Document, db: Session, exam: Exam, versions: list[ExamVersion]):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"MA TRẬN ĐỀ — {exam.title}")
    r.bold = True
    r.font.size = Pt(14)
    doc.add_paragraph("")
    if exam.template_id:
        sections = (db.query(TemplateSection).filter_by(template_id=exam.template_id)
                    .order_by(TemplateSection.section_order).all())
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Phần", "Dạng bài", "Số câu"
        for s in sections:
            row = table.add_row().cells
            row[0].text, row[1].text, row[2].text = s.name, s.question_type, str(s.question_count)
    doc.add_paragraph("")
    p = doc.add_paragraph()
    r = p.add_run("Cân bằng đáp án theo mã đề")
    r.bold = True
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text, hdr[4].text = "Mã đề", "A", "B", "C", "D"
    for v in versions:
        row = table.add_row().cells
        row[0].text = v.code
        for j, label in enumerate("ABCD", 1):
            row[j].text = str((v.answer_balance or {}).get(label, 0))
