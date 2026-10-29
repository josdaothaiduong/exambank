"""EB-401..EB-405 — Template hệ thống, sinh đề từ kho approved, hoán vị + cân bằng đáp án,
lỗi thiếu câu 422 rõ ràng, export Word 3 bản + bản in HTML."""
import io
import uuid
from collections import Counter

import pytest
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.db import Base
from app.examgen.export_docx import build_docx
from app.examgen.export_html import build_print_html
from app.models import (
    AISolution, Exam, ExamItem, ExamTemplate, ExamVersion, ImportJob, Option,
    PassageGroup, QuestionItem, QuestionStatus, QuestionTag, SourceExam, Tag,
    TemplateSection, User, Workspace,
)
from app.routers.exams import GenerateIn, generate, list_templates
from app.seed import seed_taxonomy, seed_templates

TABLES = [User.__table__, Workspace.__table__, SourceExam.__table__, PassageGroup.__table__,
          Tag.__table__, QuestionItem.__table__, Option.__table__, QuestionTag.__table__,
          ImportJob.__table__, AISolution.__table__, ExamTemplate.__table__,
          TemplateSection.__table__, Exam.__table__, ExamVersion.__table__, ExamItem.__table__]

DIFFS = ["NB", "TH", "VD", "VDC"]


@pytest.fixture(scope="module")
def ctx():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine, tables=TABLES)
    db = sessionmaker(bind=engine)()
    seed_taxonomy(db)
    seed_templates(db)
    user = User(email="gen@x.edu", full_name="Gen", password_hash="x", role="teacher")
    db.add(user); db.flush()
    ws = Workspace(owner_id=user.id, name="WS Gen")
    db.add(ws); db.flush()
    exam_src = SourceExam(workspace_id=ws.id, created_by=user.id, title="Nguồn", grade=12, exam_kind="bank")
    db.add(exam_src); db.flush()
    counter = {"h": 0}

    def add_q(qtype, diff, passage=None, approved=True):
        counter["h"] += 1
        q = QuestionItem(workspace_id=ws.id, source_exam_id=exam_src.id, created_by=user.id,
                         order_in_source=counter["h"], question_type=qtype,
                         stem_text=f"Câu {qtype} số {counter['h']} ______ kiểm tra.",
                         stem_html=f"Câu {qtype} số {counter['h']} ______ kiểm tra.",
                         difficulty=diff, status=QuestionStatus.approved if approved else QuestionStatus.pending_review,
                         dedup_hash=f"gen-{counter['h']}", answer_source="teacher_manual",
                         passage_id=passage.id if passage else None)
        db.add(q); db.flush()
        opts = []
        for j, l in enumerate("ABCD"):
            o = Option(question_id=q.id, label=l, position=j + 1,
                       content_text=f"Phương án {l} của câu {counter['h']}",
                       content_html=f"Phương án {l} của câu {counter['h']}")
            db.add(o); opts.append(o)
        db.flush()
        q.correct_option_id = opts[counter["h"] % 4].id
        db.flush()
        return q

    # Kho: ngữ pháp 12, từ vựng 12 (đủ mọi mức khó), cloze: 1 đoạn 5 câu, đọc hiểu: 1 đoạn 8 câu
    for i in range(12):
        add_q("grammar_mc", DIFFS[i % 4])
        add_q("vocabulary_mc", DIFFS[i % 4])
    pg = PassageGroup(source_exam_id=exam_src.id, workspace_id=ws.id, passage_type="cloze",
                      content_text="Đoạn điền khuyết về thành phố tương lai.",
                      content_html="<p>Đoạn điền khuyết về thành phố tương lai.</p>",
                      title="Đọc đoạn văn và chọn đáp án đúng.")
    db.add(pg); db.flush()
    for i in range(5):
        add_q("cloze", DIFFS[i % 4], passage=pg)
    pg2 = PassageGroup(source_exam_id=exam_src.id, workspace_id=ws.id, passage_type="reading",
                       content_text="Đoạn đọc hiểu dài về giao thông thông minh.",
                       content_html="<p>Đoạn đọc hiểu dài về giao thông thông minh.</p>",
                       title="Đọc đoạn văn và trả lời câu hỏi.")
    db.add(pg2); db.flush()
    for i in range(8):
        add_q("reading_comprehension", DIFFS[i % 4], passage=pg2)
    for qt in ("pronunciation", "stress", "synonym", "antonym", "sentence_arrangement"):
        for i in range(3):
            add_q(qt, DIFFS[i % 4])
    # 3 câu CHƯA duyệt — tuyệt đối không được lọt vào đề sinh ra
    for i in range(3):
        add_q("grammar_mc", "NB", approved=False)
    db.commit()
    tpl15 = db.query(ExamTemplate).filter(ExamTemplate.name.like("%15 phút%")).one()
    return {"db": db, "user": user, "ws": ws, "tpl15": tpl15}


def test_templates_seeded(ctx):
    out = list_templates(workspace_id=ctx["ws"].id, db=ctx["db"], user=ctx["user"])
    names = {t["name"]: t for t in out["items"]}
    assert len(names) == 3
    assert names["Đề thi thử TN THPT — Tiếng Anh (50 câu)"]["total_questions"] == 50
    assert sum(s["question_count"] for s in names["Đề thi thử TN THPT — Tiếng Anh (50 câu)"]["sections"]) == 50
    assert sum(s["question_count"] for s in names["Kiểm tra 45 phút — Tiếng Anh (40 câu)"]["sections"]) == 40
    assert sum(s["question_count"] for s in names["Kiểm tra 15 phút — Tiếng Anh (10 câu)"]["sections"]) == 10


def test_generate_15min_multi_version_balanced(ctx):
    db = ctx["db"]
    res = generate(GenerateIn(workspace_id=ctx["ws"].id, template_id=ctx["tpl15"].id,
                              n_versions=4, seed=42),
                   db=db, user=ctx["user"])
    assert len(res["versions"]) == 4
    total_qids = set()
    for v in res["versions"]:
        assert v["question_count"] == 10
        balance = v["answer_balance"]
        assert max(balance.values()) - min(balance.values()) <= 1, balance
    detail_items = None
    from app.routers.exams import exam_detail
    detail = exam_detail(uuid.UUID(res["id"]), db=db, user=ctx["user"])
    v1 = {it["question_id"]: it for it in detail["versions"][0]["items"]}
    total_qids = set(v1)
    assert len(total_qids) == 10
    # Mọi mã đề dùng cùng 10 câu, chỉ khác thứ tự + nhãn phương án
    for v in detail["versions"][1:]:
        assert {it["question_id"] for it in v["items"]} == total_qids
    # Đáp án sau hoán vị vẫn trỏ đúng NỘI DUNG đúng
    for v in detail["versions"]:
        for it in v["items"]:
            orig = v1[it["question_id"]]
            orig_correct_text = next(o["text"] for o in orig["options"] if o["label"] == orig["correct_label"])
            new_correct_text = next(o["text"] for o in it["options"] if o["label"] == it["correct_label"])
            assert orig_correct_text == new_correct_text


def test_generate_45min_passage_groups_intact(ctx):
    db = ctx["db"]
    tpl45 = db.query(ExamTemplate).filter(ExamTemplate.name.like("%45 phút%")).one()
    res = generate(GenerateIn(workspace_id=ctx["ws"].id, template_id=tpl45.id,
                              n_versions=2, seed=7),
                   db=db, user=ctx["user"])
    assert res["versions"][0]["question_count"] == 40
    from app.routers.exams import exam_detail
    detail = exam_detail(uuid.UUID(res["id"]), db=db, user=ctx["user"])
    items = detail["versions"][0]["items"]
    cloze = [it for it in items if it["question_type"] == "cloze"]
    reading = [it for it in items if it["question_type"] == "reading_comprehension"]
    assert len(cloze) == 5 and len({it["passage_id"] for it in cloze}) == 1
    assert len(reading) == 8 and len({it["passage_id"] for it in reading}) == 1
    # Nhóm đoạn phải LIỀN NHAU trong đề (không bị xé lẻ bởi hoán vị)
    for group in (cloze, reading):
        pos = sorted(it["position"] for it in group)
        assert pos == list(range(pos[0], pos[0] + len(pos)))


def test_shortfall_422_clear(ctx):
    db = ctx["db"]
    tpl50 = db.query(ExamTemplate).filter(ExamTemplate.name.like("%TN THPT%")).one()
    from fastapi import HTTPException
    with pytest.raises(HTTPException) as exc:
        generate(GenerateIn(workspace_id=ctx["ws"].id, template_id=tpl50.id, n_versions=1),
                 db=db, user=ctx["user"])
    assert exc.value.status_code == 422
    detail = exc.value.detail
    sections = {s["section"] for s in detail["shortfalls"]}
    assert "Tìm lỗi sai" in sections and "Viết lại câu" in sections and "Đọc hiểu" in sections
    assert all(s["available"] < s["needed"] for s in detail["shortfalls"])
    tim_loi = next(s for s in detail["shortfalls"] if s["section"] == "Tìm lỗi sai")
    assert tim_loi == {"section": "Tìm lỗi sai", "needed": 2, "available": 0}


def test_export_docx_and_print(ctx):
    db = ctx["db"]
    exam = db.query(Exam).first()
    for kind in ("student", "teacher", "matrix"):
        data = build_docx(db, exam, kind=kind)
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Mã đề" in text or "MA TRẬN" in text
        if kind == "teacher":
            assert "BẢNG ĐÁP ÁN" in text or any("BẢNG ĐÁP ÁN" in c.text for t in doc.tables for r in t.rows for c in r.cells)
    version = db.query(ExamVersion).filter_by(exam_id=exam.id).first()
    html = build_print_html(db, exam, version, teacher=False)
    assert "Câu 1:" in html and "Mã đề" in html
    assert "BẢNG ĐÁP ÁN" not in html  # bản học sinh không lộ đáp án
